from django.utils import timezone
from datetime import timedelta, date, datetime
from decimal import Decimal
from .models import AllowanceType, DeductionType, PayrollPeriod, Payslip, PayslipItem
from employee.models import Employee
from institution.models import (
    Institution,
    InstitutionBankAccount,
    InstitutionBankType,
    Department,
)
from django.shortcuts import get_object_or_404
from io import BytesIO
import io
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter
from .models import Payslip, PayrollPeriod

from datetime import date
from employee.models import Employee, EmployeeAttendance
from leave_mgt.models import LeaveApplication
from settings.models import SystemDay
from recruitment.models import JobPosition
from collections import defaultdict
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from decimal import Decimal
from django.db.models import Prefetch
import django.db.models as models
from datetime import timedelta
from dateutil.relativedelta import relativedelta
from django.utils import timezone
from io import BytesIO
from weasyprint import HTML
from decimal import Decimal
from payroll.models import AllowanceType, DeductionType



DAY_CODE_TO_WEEKDAY = {
    "MON": 0,
    "TUE": 1,
    "WED": 2,
    "THU": 3,
    "FRI": 4,
    "SAT": 5,
    "SUN": 6,
}


def get_employee_institution_working_days(employee_id):

    employee = Employee.objects.get(id=employee_id)

    if hasattr(employee, "custom_working_days"):
        working_days = employee.custom_working_days.days.values_list(
            "day_code", flat=True
        )
    else:
        institution = employee.payroll_branch.institution
        working_days = institution.working_days.days.values_list("day_code", flat=True)

    weekday_integers = {
        DAY_CODE_TO_WEEKDAY.get(day_code.upper())
        for day_code in working_days
        if day_code.upper() in DAY_CODE_TO_WEEKDAY
    }

    if not weekday_integers:
        raise ValueError("Working days are empty or invalid for employee/institution.")

    return weekday_integers


def get_employee_attendance_status_for_date(employee_id, target_date: date) -> str:
    try:
        employee = Employee.objects.get(id=employee_id)
    except Employee.DoesNotExist:
        raise Employee.DoesNotExist(f"Employee with ID {employee_id} does not exist.")

    weekday = target_date.weekday()
    working_weekdays = get_employee_institution_working_days(employee_id)

    if weekday not in working_weekdays:
        return "N-W-D"

    leave = LeaveApplication.objects.filter(
        employee=employee,
        status="approved",
        start_date__lte=target_date,
        end_date__gte=target_date,
    ).first()

    if leave:
        category = leave.leave_type.category
        return {
            "annual": "A-L",
            "sick": "S-L",
            "maternity": "M-L",
            "paternity": "P-L",
            "compassionate": "C-L",
            "study": "Sty-L",
            "unpaid": "UN-P-L",
        }.get(category, "L")

    attendance = EmployeeAttendance.objects.filter(
        employee=employee,
        status="approved",
        date=target_date,
    ).first()

    if attendance:
        if not attendance.check_in_time:
            return "absent"

        branch_open_time = getattr(employee.payroll_branch, "branch_opening_time", None)

        if not branch_open_time:
            raise ValueError("Branch opening time is not set for this employee.")

        if attendance.check_in_time <= branch_open_time:
            return "P-on-T"
        else:
            return "P-past-T"

    return "absent"


def generate_eft_excel(payroll_period_id: int, paying_account_id: int):
    try:
        payroll_period = PayrollPeriod.objects.get(id=payroll_period_id)
    except PayrollPeriod.DoesNotExist:
        raise ValueError(f"PayrollPeriod with ID {payroll_period_id} not found.")

    chosen_paying_account = InstitutionBankAccount.objects.filter(
        id=paying_account_id, institution_bank__institution=payroll_period.institution
    ).first()

    if not chosen_paying_account:
        raise ValueError(
            "No institution bank account found for the payroll period's institution."
        )

    wb = Workbook()
    ws = wb.active
    ws.title = "EFT Upload"

    bold_font = Font(bold=True)
    center_align = Alignment(horizontal="center", vertical="center")
    left_align = Alignment(horizontal="left", vertical="center")
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )
    light_blue_fill = PatternFill(
        start_color="D9E1F2", end_color="D9E1F2", fill_type="solid"
    )
    dark_blue_fill = PatternFill(
        start_color="4472C4", end_color="4472C4", fill_type="solid"
    )
    white_font = Font(color="FFFFFF", bold=True)
    red_font = Font(color="FF0000")

    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 12
    ws.column_dimensions["C"].width = 12
    ws.column_dimensions["D"].width = 18
    ws.column_dimensions["E"].width = 15
    ws.column_dimensions["F"].width = 20
    ws.column_dimensions["G"].width = 18
    ws.column_dimensions["H"].width = 18

    ws.merge_cells("D1:F1")
    ws["D1"] = "BULK EFT UPLOAD TEMPLATE"
    ws["D1"].font = bold_font
    ws["D1"].alignment = center_align

    ws["G1"] = f"Date {date.today().strftime('%d-%b-%y')}"

    ws["G1"].alignment = Alignment(horizontal="right")

    ws["A3"] = "SALARY PERIOD"
    ws["A3"].font = bold_font
    ws["B3"] = payroll_period.start_date.strftime("%m/%Y")
    ws["E3"] = "Chq Amount"
    ws["F3"] = "-"

    ws["E4"] = "Item Count"
    ws["F4"] = 0

    ws["A5"] = "DEBIT ACCOUNT*"
    ws["A5"].font = bold_font
    ws["B5"] = chosen_paying_account.account_number if chosen_paying_account else ""
    ws["B5"].fill = light_blue_fill

    ws["E5"] = "Account Name*"
    ws["E5"].font = bold_font
    ws["F5"] = chosen_paying_account.account_name if chosen_paying_account else ""
    ws["F5"].fill = light_blue_fill

    headers = [
        "DR ACCOUNT NO.",
        "BNKCODE",
        "BRCODE",
        "CR ACCOUNT*",
        "AMOUNT*",
        "BENEFICIARY NAME*",
        "DR ACCOUNT NAME",
        "BANK NAME*",
    ]
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=9, column=col_num, value=header)
        cell.font = bold_font
        cell.alignment = center_align
        cell.border = thin_border

    payslips = (
        Payslip.objects.filter(
            payroll_period=payroll_period,
            employee__payroll_branch__paying_bank_account=chosen_paying_account,
        )
        .select_related(
            "employee__user",
            "employee__payroll_branch__institution",
        )
        .prefetch_related("employee__payroll_branch__institution__banks")
    )

    row_num = 10
    total_amount = Decimal("0.00")
    item_count = 0

    for payslip in payslips:
        employee = payslip.employee
        beneficiary_name = f"{employee.user.fullname}" if employee.user else "N/A"

        bank_code = (
            chosen_paying_account.institution_bank.bank_code
            if chosen_paying_account.institution_bank
            else ""
        )
        br_code = (
            chosen_paying_account.institution_bank.br_code
            if chosen_paying_account.institution_bank
            else ""
        )

        data_row = [
            (chosen_paying_account.account_number if chosen_paying_account else ""),
            bank_code,
            br_code,
            employee.bank_account_number,
            float(payslip.net_salary),
            beneficiary_name,
            (chosen_paying_account.account_name if chosen_paying_account else ""),
            employee.bank,
        ]

        for col_num, cell_value in enumerate(data_row, 1):
            cell = ws.cell(row=row_num, column=col_num, value=cell_value)
            cell.border = thin_border
            if col_num in [
                2,
                4,
                5,
                6,
                8,
            ]:
                cell.fill = light_blue_fill
            if col_num == 5:
                cell.alignment = Alignment(horizontal="right")
            else:
                cell.alignment = left_align

        total_amount += payslip.net_salary
        item_count += 1
        row_num += 1

    ws["F4"] = item_count
    ws["F3"] = float(total_amount)

    excel_file = io.BytesIO()
    wb.save(excel_file)
    excel_file.seek(0)

    return excel_file


class PayrollProcessor:
    """
    Utility class to handle payroll processing operations
    """

    @staticmethod
    def create_monthly_period(year, month):
        """
        Create a monthly payroll period
        """
        from calendar import monthrange

        start_date = datetime(year, month, 1).date()
        last_day = monthrange(year, month)[1]
        end_date = datetime(year, month, last_day).date()

        # Set pay_date to last working day or a few days after month end
        pay_date = end_date + timedelta(days=3)

        period_name = f"{start_date.strftime('%B %Y')}"

        period, created = PayrollPeriod.objects.get_or_create(
            name=period_name,
            start_date=start_date,
            defaults={
                "end_date": end_date,
                "pay_date": pay_date,
            },
        )
        return period, created

    @staticmethod
    def generate_payslips_for_period(payroll_period, employee_ids=None):
        """
        Generate payslips for a given payroll period and mark the period as processed.
        """
        if isinstance(payroll_period, int):
            payroll_period = get_object_or_404(PayrollPeriod, id=payroll_period)

        employees = Employee.objects.filter(
            id__in=employee_ids,
            is_active=True,
        )

        created_payslips = []

        for employee in employees:
            payslip, created = Payslip.objects.get_or_create(
                employee=employee,
                payroll_period=payroll_period,
                defaults={
                    "basic_salary": employee.salary or 0,
                },
            )

            if created or not payslip.is_paid:
                PayslipItem.generate_payslip_items(payslip)
                payslip.calculate_totals()
                # PayrollProcessor.generate_payslip_items(payslip)
                created_payslips.append(payslip)

        if created_payslips:
            payroll_period.is_processed = True
            payroll_period.save()

        return created_payslips

    # @staticmethod
    # def generate_payslip_items(payslip):
    #     """
    #     Generate detailed payslip items for allowances and deductions
    #     """
    #     # Clear existing items
    #     payslip.items.all().delete()

    #     items_to_create = []

    #     # Add allowances
    #     for allowance in payslip.employee.allowances.filter(is_active=True):
    #         if allowance.effective_from <= payslip.payroll_period.end_date and (
    #             not allowance.effective_to
    #             or allowance.effective_to >= payslip.payroll_period.start_date
    #         ):

    #             amount = allowance.get_calculated_amount()
    #             items_to_create.append(
    #                 PayslipItem(
    #                     payslip=payslip,
    #                     item_type="allowance",
    #                     name=allowance.allowance_type.name,
    #                     amount=amount,
    #                     description=f"{allowance.calculation_method}: {allowance.amount if allowance.calculation_method == 'fixed' else f'{allowance.percentage}%'}",
    #                 )
    #             )

    #     # Add deductions
    #     for deduction in payslip.employee.deductions.filter(is_active=True):
    #         if deduction.effective_from <= payslip.payroll_period.end_date and (
    #             not deduction.effective_to
    #             or deduction.effective_to >= payslip.payroll_period.start_date
    #         ):

    #             amount = deduction.get_calculated_amount()
    #             items_to_create.append(
    #                 PayslipItem(
    #                     payslip=payslip,
    #                     item_type="deduction",
    #                     name=deduction.deduction_type.name,
    #                     amount=amount,
    #                     description=f"{deduction.calculation_method}: {deduction.amount if deduction.calculation_method == 'fixed' else f'{deduction.percentage}%'}",
    #                 )
    #             )

    #     # Add overtime if applicable
    #     # if payslip.overtime_amount > 0:
    #     #     items_to_create.append(
    #     #         PayslipItem(
    #     #             payslip=payslip,
    #     #             item_type='overtime',
    #     #             name='Overtime Pay',
    #     #             amount=payslip.overtime_amount,
    #     #             description=f"{payslip.overtime_hours} hours @ {payslip.overtime_rate} per hour"
    #     #         )
    #     #     )

    #     # Bulk create items
    #     PayslipItem.objects.bulk_create(items_to_create)

    @staticmethod
    def setup_default_payroll_types_for_institution(institution):
        """
        Create default allowance and deduction types for a specific institution
        """

        # Default allowances
        default_allowances = [
            {
                "name": "Housing Allowance",
                "description": "Monthly housing allowance",
                "is_taxable": True,
            },
            {
                "name": "Transport Allowance",
                "description": "Monthly transport allowance",
                "is_taxable": True,
            },
            {
                "name": "Medical Allowance",
                "description": "Monthly medical allowance",
                "is_taxable": False,
            },
            {
                "name": "Lunch Allowance",
                "description": "Daily lunch allowance",
                "is_taxable": True,
            },
            {
                "name": "Bonus",
                "description": "Performance or annual bonus",
                "is_taxable": True,
            },
            {
                "name": "Overtime Allowance",
                "description": "Overtime payment allowance",
                "is_taxable": True,
            },
            {
                "name": "Communication Allowance",
                "description": "Monthly communication allowance",
                "is_taxable": True,
            },
        ]

        # Default deductions
        default_deductions = [
            {
                "name": "Health Insurance",
                "description": "Monthly health insurance premium",
                "is_mandatory": False,
            },
            {
                "name": "Loan Repayment",
                "description": "Monthly loan repayment",
                "is_mandatory": False,
            },
            {
                "name": "Union Dues",
                "description": "Monthly union membership fees",
                "is_mandatory": False,
            },
            {
                "name": "Advance Salary",
                "description": "Salary advance repayment",
                "is_mandatory": False,
            },
        ]

        created_allowances = []
        created_deductions = []

        # Create allowance types
        for allowance_data in default_allowances:
            allowance_type, created = AllowanceType.objects.get_or_create(
                institution=institution,
                name=allowance_data["name"],
                defaults={
                    "description": allowance_data["description"],
                    "is_taxable": allowance_data["is_taxable"],
                    "is_active": True,
                },
            )
            if created:
                created_allowances.append(allowance_type)

        # Create deduction types
        for deduction_data in default_deductions:
            deduction_type, created = DeductionType.objects.get_or_create(
                institution=institution,
                name=deduction_data["name"],
                defaults={
                    "description": deduction_data["description"],
                    "is_mandatory": deduction_data["is_mandatory"],
                    "is_active": True,
                },
            )
            if created:
                created_deductions.append(deduction_type)

        return {
            "allowances_created": len(created_allowances),
            "deductions_created": len(created_deductions),
            "allowance_types": created_allowances,
            "deduction_types": created_deductions,
        }

    @staticmethod  # Fixed: Added @staticmethod decorator
    def create_institution_with_defaults(**kwargs):
        """
        Helper function to create institution with default payroll types
        """
        institution = Institution.objects.create(**kwargs)
        setup_result = PayrollProcessor.setup_default_payroll_types_for_institution(
            institution
        )

        return institution, setup_result

    @staticmethod
    def get_employee_payroll_summary(employee, year=None):
        """
        Get annual payroll summary for an employee
        """
        if not year:
            year = timezone.now().year

        payslips = Payslip.objects.filter(
            employee=employee, payroll_period__start_date__year=year
        ).order_by("payroll_period__start_date")

        summary = {
            "employee": employee,
            "year": year,
            "total_gross": sum(p.gross_salary for p in payslips),
            "total_net": sum(p.net_salary for p in payslips),
            "total_allowances": sum(p.total_allowances for p in payslips),
            "total_deductions": sum(p.total_deductions for p in payslips),
            "payslips_count": payslips.count(),
            "payslips": payslips,
        }

        return summary


# Utility functions that can be used in management commands or views
def create_current_month_period():
    """Create payroll period for current month"""
    now = timezone.now()
    return PayrollProcessor.create_monthly_period(now.year, now.month)


def process_monthly_payroll(year=None, month=None, employee_ids=None):
    """
    Complete monthly payroll processing
    """
    if not year or not month:
        now = timezone.now()
        year = year or now.year
        month = month or now.month

    # Create period
    period, created = PayrollProcessor.create_monthly_period(year, month)

    # Generate payslips
    payslips = PayrollProcessor.generate_payslips_for_period(period, employee_ids)

    return {
        "period": period,
        "period_created": created,
        "payslips_generated": len(payslips),
        "payslips": payslips,
    }


FILL_PRESENT = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
FILL_LATE = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
FILL_ABSENT = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
FILL_ANNUAL = PatternFill(start_color="BDD7EE", end_color="BDD7EE", fill_type="solid")
FILL_SICK = PatternFill(start_color="D9D2E9", end_color="D9D2E9", fill_type="solid")
FILL_OTHER_LEAVE = PatternFill(
    start_color="E0E0E0", end_color="E0E0E0", fill_type="solid"
)
FILL_NON_WORKING = PatternFill(
    start_color="F2F2F2", end_color="F2F2F2", fill_type="solid"
)
FILL_ERROR = PatternFill(start_color="A93226", end_color="A93226", fill_type="solid")

FONT_WHITE = Font(color="FFFFFF")
FONT_GRAY = Font(color="808080")
BOLD_FONT = Font(bold=True)


def generate_attendance_excel(
    start_date: date, end_date: date, context: dict
) -> BytesIO:
    employee_qs = Employee.objects.filter(is_active=True).select_related("user")

    target_ids = context.get("target_employees")
    target_departments = context.get("target_departments")
    target_positions = context.get("target_job_positions")

    if target_ids:
        employee_qs = employee_qs.filter(id__in=target_ids)
    elif target_departments:
        employee_qs = employee_qs.filter(department__in=target_departments)
    elif target_positions:
        employee_qs = employee_qs.filter(position__in=target_positions)

    employees = employee_qs.order_by("user__fullname")
    if not employees.exists():
        raise ValueError("No employees found for the provided filters.")

    num_days = (end_date - start_date).days + 1
    date_list = [start_date + timedelta(days=i) for i in range(num_days)]

    wb = Workbook()
    ws = wb.active
    ws.title = "Attendance Report"

    center_align = Alignment(horizontal="center", vertical="center")
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )
    header_fill = PatternFill(
        start_color="D9E1F2", end_color="D9E1F2", fill_type="solid"
    )

    legend_data = [
        ("P-on-T", "Present on Time", FILL_PRESENT),
        ("P-past-T", "Present but Late", FILL_LATE),
        ("absent", "Absent (no check-in)", FILL_ABSENT),
        ("A-L", "Annual Leave", FILL_ANNUAL),
        ("S-L", "Sick Leave", FILL_SICK),
        ("M-L", "Maternity Leave", FILL_OTHER_LEAVE),
        ("P-L", "Paternity Leave", FILL_OTHER_LEAVE),
        ("C-L", "Compassionate Leave", FILL_OTHER_LEAVE),
        ("Sty-L", "Study Leave", FILL_OTHER_LEAVE),
        ("UN-P-L", "Unpaid Leave", FILL_OTHER_LEAVE),
        ("N-W-D", "Not a Working Day", FILL_NON_WORKING),
        ("ERR", "Error fetching status", FILL_ERROR),
    ]

    for idx, (code, description, fill) in enumerate(legend_data, start=1):
        ws.cell(row=idx, column=1, value=code).fill = fill
        ws.cell(row=idx, column=1).font = BOLD_FONT
        ws.cell(row=idx, column=2, value=description).alignment = Alignment(
            wrap_text=True
        )
        if code == "ERR":
            ws.cell(row=idx, column=1).font = FONT_WHITE

    start_data_row = len(legend_data) + 2

    headers = [
        "Employee ID",
        "Full Name",
        "Present",
        "Absent",
        "Late",
        "Leave",
        "Total Days",
    ]
    for col_num, header in enumerate(headers, start=1):
        cell = ws.cell(row=start_data_row, column=col_num, value=header)
        cell.font = BOLD_FONT
        cell.fill = header_fill
        cell.alignment = center_align
        cell.border = thin_border
        ws.column_dimensions[cell.column_letter].width = 15

    for col_num, current_date in enumerate(date_list, start=len(headers) + 1):
        cell = ws.cell(
            row=start_data_row, column=col_num, value=current_date.strftime("%d-%b")
        )
        cell.font = BOLD_FONT
        cell.fill = header_fill
        cell.alignment = center_align
        cell.border = thin_border
        ws.column_dimensions[cell.column_letter].width = 12

    for row_offset, employee in enumerate(employees, start=start_data_row + 1):
        present_count = 0
        absent_count = 0
        late_count = 0
        leave_count = 0
        total_working_days = 0

        ws.cell(row=row_offset, column=1, value=employee.id).border = thin_border
        full_name = employee.user.fullname if employee.user else "N/A"
        ws.cell(row=row_offset, column=2, value=full_name).border = thin_border

        for col_offset, current_date in enumerate(date_list, start=len(headers) + 1):
            try:
                # status = get_employee_attendance_status_for_date(
                #     employee.id, current_date
                # )
                status = []
            except Exception:
                status = "ERR"

            if status == "P-on-T":
                present_count += 1
                total_working_days += 1
            elif status == "P-past-T":
                late_count += 1
                total_working_days += 1
            elif status == "absent":
                absent_count += 1
                total_working_days += 1
            elif status in ["A-L", "S-L", "M-L", "P-L", "C-L", "Sty-L", "UN-P-L"]:
                leave_count += 1
                total_working_days += 1

            cell = ws.cell(row=row_offset, column=col_offset, value=status)
            cell.border = thin_border
            cell.alignment = center_align

            if status == "P-on-T":
                cell.fill = FILL_PRESENT
            elif status == "P-past-T":
                cell.fill = FILL_LATE
            elif status == "absent":
                cell.fill = FILL_ABSENT
            elif status == "A-L":
                cell.fill = FILL_ANNUAL
            elif status == "S-L":
                cell.fill = FILL_SICK
            elif status in ["M-L", "P-L", "C-L", "Sty-L", "UN-P-L"]:
                cell.fill = FILL_OTHER_LEAVE
            elif status == "N-W-D":
                cell.fill = FILL_NON_WORKING
                cell.font = FONT_GRAY
            elif status == "ERR":
                cell.fill = FILL_ERROR
                cell.font = FONT_WHITE

        summary = [
            present_count,
            absent_count,
            late_count,
            leave_count,
            total_working_days,
        ]
        for idx, value in enumerate(summary, start=3):
            cell = ws.cell(row=row_offset, column=idx, value=value)
            cell.border = thin_border

    ws.freeze_panes = ws.cell(row=start_data_row + 1, column=len(headers) + 1)

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output


def generate_allpayslips_excel(payroll_period_id: int) -> BytesIO:
    payslips = Payslip.objects.filter(
        payroll_period_id=payroll_period_id
    ).prefetch_related(
        Prefetch("items"),
        "employee__user",
    )

    if not payslips.exists():
        raise ValueError("No payslips found for the given payroll period.")

    allowance_names = set()
    deduction_names = set()

    for payslip in payslips:
        for item in payslip.items.all():
            if item.item_type == "allowance":
                allowance_names.add(item.name)
            elif item.item_type == "deduction":
                deduction_names.add(item.name)

    allowance_names = sorted(allowance_names)
    deduction_names = sorted(deduction_names)

    base_headers = [
        "Full Name",
        "Basic Salary",
        "Gross Salary",
        "Net Salary",
    ]
    deduction_headers = [f"Deduction - {name}" for name in deduction_names]
    allowance_headers = [f"Allowance - {name}" for name in allowance_names]
    final_headers = base_headers + deduction_headers + allowance_headers

    wb = Workbook()
    ws = wb.active
    ws.title = "Payslips"

    # Styling
    bold_font = Font(bold=True)
    center_align = Alignment(horizontal="center")
    header_fill = PatternFill(
        start_color="D9E1F2", end_color="D9E1F2", fill_type="solid"
    )
    total_fill = PatternFill(
        start_color="E2EFDA", end_color="E2EFDA", fill_type="solid"
    )
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )

    totals = defaultdict(Decimal)

    # Write header in row 2
    for col_idx, header in enumerate(final_headers, start=1):
        cell = ws.cell(row=2, column=col_idx, value=header)
        cell.font = bold_font
        cell.alignment = center_align
        cell.fill = header_fill
        cell.border = thin_border
        ws.column_dimensions[cell.column_letter].width = max(15, len(header) + 2)

    for row_idx, payslip in enumerate(payslips, start=3):
        employee = payslip.employee
        employee_name = employee.user.fullname if employee.user else "N/A"

        base_row = [
            employee_name,
            float(payslip.basic_salary),
            float(payslip.gross_salary),
            float(payslip.net_salary),
        ]

        deduction_map = defaultdict(lambda: Decimal("0.00"))
        allowance_map = defaultdict(lambda: Decimal("0.00"))

        for item in payslip.items.all():
            if item.item_type == "deduction":
                deduction_map[item.name] += item.amount
            elif item.item_type == "allowance":
                allowance_map[item.name] += item.amount

        row = base_row
        for name in deduction_names:
            row.append(float(deduction_map.get(name, 0.00)))

        for name in allowance_names:
            row.append(float(allowance_map.get(name, 0.00)))

        for col_idx, value in enumerate(row, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.alignment = center_align
            cell.border = thin_border

            if isinstance(value, (int, float)):
                totals[col_idx] += Decimal(str(value))

    for col_idx in range(1, len(final_headers) + 1):
        value = totals.get(col_idx)
        if value is not None:
            cell = ws.cell(row=1, column=col_idx, value=float(value))
            cell.font = bold_font
            cell.alignment = center_align
            cell.fill = total_fill
            cell.border = thin_border

    ws.cell(row=1, column=1, value="TOTALS").font = bold_font

    ws.freeze_panes = ws["A3"]

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output

#TODO: Ensure that decimals are not just formatted since there is a possibility of them being null and we cause issues
# def generate_payslip_pdf(payslip):
#     """
#     Generates a structured payslip PDF with a detailed breakdown of the employee's
#     specific allowances, deductions, and taxes, only including items applicable to the employee.
#     """
#     from decimal import Decimal
#     from collections import defaultdict

#     # Safely get institution and company name
#     institution = None
#     company_name = 'N/A'
#     currency = 'UGX'
#     if payslip.employee and payslip.employee.department and payslip.employee.department.institution:
#         institution = payslip.employee.department.institution
#         company_name = institution.institution_name
#         # Assuming the Institution model has a 'currency' field
#         if hasattr(institution, 'currency') and institution.currency:
#             currency = institution.currency

#     # Get the employee's specific allowances, deductions, and taxes
#     employee = payslip.employee


#     # Fetch relevant allowances
#     allowances = (
#         employee.allowances.filter(is_active=True)
#         .filter(
#             effective_from__lte=payslip.payroll_period.end_date,
#         )
#         .filter(
#             models.Q(effective_to__gte=payslip.payroll_period.start_date)
#             | models.Q(effective_to__isnull=True)
#         )
#     )
#     allowances_data = []
#     for allowance in allowances:
#         recurrence = allowance.get_recurrence_count(payslip.payroll_period)
#         amount = allowance.get_calculated_amount() * recurrence
#         allowances_data.append({'name': allowance.allowance_type.name, 'amount': amount})

#     # Fetch regular deductions (excluding attendance penalties)
#     deductions = (
#         employee.deductions.filter(is_active=True)
#         .exclude(deduction_type__attendance_penalty_for__in=["late", "absentism"])
#         .filter(
#             effective_from__lte=payslip.payroll_period.end_date,
#         )
#         .filter(
#             models.Q(effective_to__gte=payslip.payroll_period.start_date)
#             | models.Q(effective_to__isnull=True)
#         )
#     )
#     regular_deductions = []
#     for deduction in deductions:
#         recurrence = deduction.get_recurrence_count(payslip.payroll_period)
#         amount = deduction.get_calculated_amount() * recurrence
#         regular_deductions.append({'name': deduction.deduction_type.name, 'amount': amount})

#     # Fetch attendance deductions
#     attendance_total, deduction_items = payslip.get_attendance_deductions()
#     attendance_ded_group = defaultdict(Decimal)
#     for item in deduction_items:
#         attendance_ded_group[item['name']] += item['amount']
#     attendance_deductions = [{'name': name, 'amount': amt} for name, amt in attendance_ded_group.items()]

#     # Fetch relevant taxes
#     taxes = employee.taxes.filter(
#         effective_from__lte=payslip.payroll_period.end_date,
#     ).filter(
#         models.Q(effective_to__gte=payslip.payroll_period.start_date)
#         | models.Q(effective_to__isnull=True)
#     )
#     taxes_data = []
#     for tax in taxes:
#         amount = tax.get_tax_amount()
#         taxes_data.append({'name': tax.institution_tax.tax_name, 'amount': amount})

#     # Combine all deductions
#     deductions_data = regular_deductions + attendance_deductions + taxes_data
#     # Sort deductions by name for better presentation
#     deductions_data.sort(key=lambda x: x['name'])

#     # Filter out zero-amount items if desired (optional; comment out if you want to include zeros)
#     allowances_data = [item for item in allowances_data if item['amount'] > 0]
#     deductions_data = [item for item in deductions_data if item['amount'] > 0]

#     # Get the month and year for the header
#     pay_period_month = payslip.payroll_period.start_date.strftime('%B %Y')

#     # Safely get employee and position details
#     employee_name = payslip.employee.user.fullname if payslip.employee and payslip.employee.user else 'N/A'
#     employee_id = payslip.employee.employee_id if payslip.employee else 'N/A'
#     department_name = payslip.employee.department.name if payslip.employee and payslip.employee.department else 'N/A'
#     position_name = payslip.employee.position.name if payslip.employee and payslip.employee.position else 'N/A'

#     html_template = f"""
#     <!DOCTYPE html>
#     <html lang="en">
#     <head>
#         <meta charset="UTF-8">
#         <title>Payslip</title>
#         <style>
#             @page {{
#                 size: A4;
#                 margin: 40px;
#             }}
#             body {{
#                 font-family: Arial, sans-serif;
#                 font-size: 10px;
#                 line-height: 1.4;
#                 color: #333;
#             }}
#             .header {{
#                 text-align: center;
#                 margin-bottom: 30px;
#             }}
#             .company-name {{
#                 font-size: 24px;
#                 font-weight: bold;
#                 color: #2c3e50;
#             }}
#             .period-details {{
#                 font-size: 12px;
#                 color: #7f8c8d;
#             }}
#             h2 {{
#                 font-size: 16px;
#                 font-weight: bold;
#                 border-bottom: 2px solid #34495e;
#                 padding-bottom: 5px;
#                 margin-top: 30px;
#                 margin-bottom: 15px;
#             }}
#             table {{
#                 width: 100%;
#                 border-collapse: collapse;
#                 margin-bottom: 20px;
#             }}
#             .info-table td {{
#                 padding: 5px;
#             }}
#             .info-table td:nth-child(odd) {{
#                 font-weight: bold;
#                 width: 30%;
#             }}
#             .financial-table {{
#                 border: 1px solid #bdc3c7;
#             }}
#             .financial-table th {{
#                 background-color: #ecf0f1;
#                 font-weight: bold;
#                 padding: 10px;
#                 text-align: left;
#                 border: 1px solid #bdc3c7;
#             }}
#             .financial-table td {{
#                 padding: 8px 10px;
#                 border: 1px solid #bdc3c7;
#             }}
#             .financial-table td:last-child {{
#                 text-align: right;
#             }}
#             .summary-table {{
#                 margin-top: 20px;
#                 width: 100%;
#             }}
#             .summary-table td {{
#                 padding: 8px 10px;
#                 font-weight: bold;
#                 border: 1px solid #bdc3c7;
#             }}
#             .summary-table tr:last-child {{
#                 font-size: 14px;
#                 background-color: #e0f2f1;
#             }}
#             .summary-table td:last-child {{
#                 text-align: right;
#             }}
#             .summary-table .label {{
#                 font-weight: normal;
#                 width: 60%;
#             }}
#             .footer {{
#                 clear: both;
#                 padding-top: 20px;
#                 text-align: center;
#                 font-style: italic;
#                 font-size: 9px;
#             }}
#         </style>
#     </head>
#     <body>
#         <div class="header">
#             <div class="company-name">{company_name}</div>
#             <div class="period-details">
#                 Payslip for {pay_period_month}
#             </div>
#             <div class="period-details">
#                 Period: {payslip.payroll_period.start_date.strftime('%d-%b-%Y')} to {payslip.payroll_period.end_date.strftime('%d-%b-%Y')}
#             </div>
#         </div>

#         <h2>Employee Information</h2>
#         <table class="info-table">
#             <tr>
#                 <td>Employee Name:</td>
#                 <td>{employee_name}</td>
#                 <td>Employee ID:</td>
#                 <td>{employee_id}</td>
#             </tr>
#             <tr>
#                 <td>Department:</td>
#                 <td>{department_name}</td>
#                 <td>Position:</td>
#                 <td>{position_name}</td>
#             </tr>
#             <tr>
#                 <td>Days Worked:</td>
#                 <td>{payslip.days_worked}</td>
#                 <td>Pay Date:</td>
#                 <td>{payslip.payroll_period.pay_date.strftime('%d-%b-%Y')}</td>
#             </tr>
#         </table>
        
#         <h2>Summary</h2>
#         <table class="summary-table">
#             <tbody>
#                 <tr>
#                     <td class="label">Gross Salary:</td>
#                     <td>{payslip.gross_salary}</td>
#                 </tr>
#                 <tr>
#                     <td class="label">Total Allowances:</td>
#                     <td>{payslip.total_allowances}</td>
#                 </tr>
#                 <tr>
#                     <td class="label">Total Deductions:</td>
#                     <td>{payslip.total_deductions}</td>
#                 </tr>
#                 <tr>
#                     <td class="label">Net Salary:</td>
#                     <td>{payslip.net_salary}</td>
#                 </tr>
#             </tbody>
#         </table>

#         <h2>Allowances</h2>
#         <table class="financial-table">
#             <thead>
#                 <tr>
#                     <th>Description</th>
#                     <th>Amount ({currency})</th>
#                 </tr>
#             </thead>
#             <tbody>
#                 <tr>
#                     <td>Basic Salary</td>
#                     <td>{payslip.basic_salary}</td>
#                 </tr>
#                 {"".join(f"""
#                 <tr>
#                     <td>{item['name']}</td>
#                     <td>{item['amount']}</td>
#                 </tr>""" for item in allowances_data)}
#             </tbody>
#         </table>

#         <h2>Deductions</h2>
#         <table class="financial-table">
#             <thead>
#                 <tr>
#                     <th>Description</th>
#                     <th>Amount ({currency})</th>
#                 </tr>
#             </thead>
#             <tbody>
#                 {"".join(f"""
#                 <tr>
#                     <td>{item['name']}</td>
#                     <td>{item['amount']}</td>
#                 </tr>""" for item in deductions_data)}
#             </tbody>
#         </table>
        
#         <div class="footer">
#             This is a computer-generated document and does not require a signature.
#         </div>
#     </body>
#     </html>
#     """
    
#     html_doc = HTML(string=html_template)
#     pdf_buffer = BytesIO()
#     html_doc.write_pdf(target=pdf_buffer)
#     pdf_buffer.seek(0)
    
#     return pdf_buffer