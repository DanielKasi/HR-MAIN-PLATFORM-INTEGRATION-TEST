from django.shortcuts import render
from rest_framework.views import APIView
from rest_framework.response import Response
from rest_framework import status
from drf_spectacular.utils import extend_schema, OpenApiParameter
from .models import DocumentType, DocumentTemplate, Document
from .serializers import (
    DocumentTypeSerializer,
    DocumentTemplateSerializer,
    GenerateDocumentResponseSerializer,
    GenerateDocumentRequestSerializer,
    DocumentContentPreviewSerializer,
    DocumentStatusUpdateSerializer,
)
from utilities.pagination import CustomPageNumberPagination
from rest_framework.parsers import MultiPartParser, FormParser
from institution.models import Institution
from audit.models import AuditLog
from django.contrib.contenttypes.models import ContentType
from onboarding.models import OnBoarding
from settings.models import SystemConfiguration
from django.shortcuts import get_object_or_404
from django.core.mail import EmailMultiAlternatives
from django.conf import settings
from django.template.loader import render_to_string
from datetime import datetime, timedelta
from bs4 import BeautifulSoup
import re
import html
from docx import Document as DocxDocument
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from django.core.files.base import ContentFile
from weasyprint import HTML
from employee.models import EmployeeContract, Employee
from docxtpl import DocxTemplate
import logging

# Set up logging
logger = logging.getLogger(__name__)


class DocumentTypeListCreateAPIView(APIView):

    @extend_schema(
        description="Retrieve a paginated list of document types for a given institution.",
        responses={200: DocumentTypeSerializer(many=True)},
    )
    def get(self, request, institution_id):
        queryset = DocumentType.objects.filter(institution_id=institution_id, is_active=True).order_by(
            "-created_at"
        )
        paginator = CustomPageNumberPagination()
        paginated_qs = paginator.paginate_queryset(queryset, request)
        serializer = DocumentTypeSerializer(paginated_qs, many=True)
        return paginator.get_paginated_response(serializer.data)

    @extend_schema(
        description="Create a new document type for the given institution. The code is auto-generated from the name.",
        request=DocumentTypeSerializer,
        responses={201: DocumentTypeSerializer},
    )
    def post(self, request, institution_id):
        data = request.data.copy()
        data["institution"] = institution_id
        serializer = DocumentTypeSerializer(data=data)
        if serializer.is_valid():
            serializer.save()
            return Response(serializer.data, status=status.HTTP_201_CREATED)
        return Response(serializer.errors, status=status.HTTP_400_BAD_REQUEST)


class DocumentTypeRetrieveUpdateDeleteAPIView(APIView):

    def get_object(self, pk):
        try:
            return DocumentType.objects.get(pk=pk)
        except DocumentType.DoesNotExist:
            return None

    @extend_schema(
        description="Retrieve a specific document type by ID",
        responses={200: DocumentTypeSerializer, 404: None},
    )
    def get(self, request, pk):
        document_type = self.get_object(pk)
        if not document_type:
            return Response(status=status.HTTP_404_NOT_FOUND)
        serializer = DocumentTypeSerializer(document_type)
        return Response(serializer.data)

    @extend_schema(
        description="Partially update a document type. The code is automatically regenerated if name changes.",
        request=DocumentTypeSerializer,
        responses={200: DocumentTypeSerializer, 404: None},
    )
    def patch(self, request, pk):
        document_type = self.get_object(pk)
        if not document_type:
            return Response(status=status.HTTP_404_NOT_FOUND)
        serializer = DocumentTypeSerializer(
            document_type, data=request.data, partial=True
        )
        if serializer.is_valid():
            serializer.save()
            return Response(serializer.data)
        return Response(serializer.errors, status=status.HTTP_400_BAD_REQUEST)

    @extend_schema(
        description="Delete a document type", responses={204: None, 404: None}
    )
    def delete(self, request, pk):
        document_type = self.get_object(pk)
        if not document_type:
            return Response(status=status.HTTP_404_NOT_FOUND)
        document_type.delete()
        return Response(status=status.HTTP_204_NO_CONTENT)


class DocumentTemplateListCreateAPIView(APIView):
    parser_classes = [MultiPartParser, FormParser]

    @extend_schema(
        request=DocumentTemplateSerializer,
        responses={201: DocumentTemplateSerializer, 400: None},
        description="Create a new document template.",
        tags=["Document Templates"],
    )
    def post(self, request, institution_id):
        try:
            institution = Institution.objects.get(pk=institution_id)
        except Institution.DoesNotExist:
            return Response(
                {"error": "Institution not found"}, status=status.HTTP_404_NOT_FOUND
            )

        # Ensure document_type belongs to the institution
        document_type_id = request.data.get("document_type")
        if document_type_id:
            try:
                document_type = DocumentType.objects.get(
                    pk=document_type_id, institution=institution
                )
            except DocumentType.DoesNotExist:
                return Response(
                    {"error": "Document type does not belong to this institution"},
                    status=status.HTTP_400_BAD_REQUEST,
                )

        serializer = DocumentTemplateSerializer(
            data=request.data,
            context={"request": request, "institution_id": institution_id},
        )
        if serializer.is_valid():
            template = serializer.save()
            # Create audit log
            AuditLog.objects.create(
                content_type=ContentType.objects.get_for_model(DocumentTemplate),
                object_id=template.pk,
                action="CREATE",
                user=request.user if request.user.is_authenticated else None,
                description=f"Created DocumentTemplate: {template.name}",
            )
            return Response(serializer.data, status=status.HTTP_201_CREATED)
        return Response(serializer.errors, status=status.HTTP_400_BAD_REQUEST)

    @extend_schema(
        responses={200: DocumentTemplateSerializer(many=True), 404: None},
        description="Retrieve a list of document templates for an institution.",
        tags=["Document Templates"],
    )
    def get(self, request, institution_id):
        try:
            Institution.objects.get(pk=institution_id)
        except Institution.DoesNotExist:
            return Response(
                {"error": "Institution not found"}, status=status.HTTP_404_NOT_FOUND
            )

        templates = DocumentTemplate.objects.filter(
            document_type__institution=institution_id
        ).order_by("-created_at")
        paginator = CustomPageNumberPagination()
        paginated_templates = paginator.paginate_queryset(templates, request)
        serializer = DocumentTemplateSerializer(
            paginated_templates, many=True, context={"request": request}
        )
        return paginator.get_paginated_response(serializer.data)


class DocumentTemplateDetailAPIView(APIView):
    parser_classes = [MultiPartParser, FormParser]

    def get_object(self, pk):
        try:
            return DocumentTemplate.objects.get(pk=pk)
        except DocumentTemplate.DoesNotExist:
            return None

    @extend_schema(
        responses={200: DocumentTemplateSerializer, 404: None},
        description="Retrieve a specific document template by ID.",
        tags=["Document Templates"],
    )
    def get(self, request, pk):
        template = self.get_object(pk)
        if not template:
            return Response(
                {"error": "Document template not found"},
                status=status.HTTP_404_NOT_FOUND,
            )
        serializer = DocumentTemplateSerializer(template, context={"request": request})
        return Response(serializer.data)

    @extend_schema(
        request=DocumentTemplateSerializer,
        responses={200: DocumentTemplateSerializer, 400: None, 404: None},
        description="Update a document template.",
        tags=["Document Templates"],
    )
    def patch(self, request, pk):
        template = self.get_object(pk)
        if not template:
            return Response(
                {"error": "Document template not found"},
                status=status.HTTP_404_NOT_FOUND,
            )

        # Capture old state for changes
        old_data = {
            "document_type": template.document_type_id,
            "name": template.name,
            "template_type": template.template_type,
            "file": template.file.name if template.file else None,
            "content": template.content,
            "placeholders": template.placeholders,
        }

        serializer = DocumentTemplateSerializer(
            template, data=request.data, partial=True, context={"request": request}
        )
        if serializer.is_valid():
            updated_template = serializer.save()
            # Calculate changes
            changes = {}
            new_data = {
                "document_type": updated_template.document_type_id,
                "name": updated_template.name,
                "template_type": updated_template.template_type,
                "file": updated_template.file.name if updated_template.file else None,
                "content": updated_template.content,
                "placeholders": updated_template.placeholders,
            }
            for field, old_value in old_data.items():
                new_value = new_data[field]
                if old_value != new_value:
                    changes[field] = {"old": old_value, "new": new_value}

            # Create audit log
            AuditLog.objects.create(
                content_type=ContentType.objects.get_for_model(DocumentTemplate),
                object_id=updated_template.pk,
                action="UPDATE",
                user=request.user if request.user.is_authenticated else None,
                changes=changes if changes else None,
                description=f"Updated DocumentTemplate: {updated_template.name}",
            )
            return Response(serializer.data)
        return Response(serializer.errors, status=status.HTTP_400_BAD_REQUEST)

    @extend_schema(
        responses={204: None, 404: None},
        description="Delete a document template.",
        tags=["Document Templates"],
    )
    def delete(self, request, pk):
        template = self.get_object(pk)
        if not template:
            return Response(
                {"error": "Document template not found"},
                status=status.HTTP_404_NOT_FOUND,
            )

        # Create audit log before deletion
        AuditLog.objects.create(
            content_type=ContentType.objects.get_for_model(DocumentTemplate),
            object_id=template.pk,
            action="DELETE",
            user=request.user if request.user.is_authenticated else None,
            description=f"Deleted DocumentTemplate: {template.name}",
        )
        template.delete()
        return Response(status=status.HTTP_204_NO_CONTENT)


class AuditLogListAPIView(APIView):
    @extend_schema(
        responses={200: None, 404: None},  # Add serializer if needed
        description="Retrieve audit logs for a specific document template.",
        tags=["Document Templates"],
    )
    def get(self, request, institution_id, template_id):
        try:
            DocumentTemplate.objects.get(pk=template_id)
        except DocumentTemplate.DoesNotExist:
            return Response(
                {"error": "Document template not found"},
                status=status.HTTP_404_NOT_FOUND,
            )

        content_type = ContentType.objects.get_for_model(DocumentTemplate)
        audit_logs = (
            AuditLog.objects.filter(content_type=content_type, object_id=template_id)
            .select_related("user")
            .order_by("-timestamp")
        )

        response_data = [
            {
                "id": log.id,
                "action": log.action,
                "user": log.user.username if log.user else None,
                "timestamp": log.timestamp.isoformat(),
                "changes": log.changes,
                "description": log.description,
            }
            for log in audit_logs
        ]
        return Response(response_data, status=status.HTTP_200_OK)


class BaseDocumentView(APIView):
    def _normalize_placeholder(self, placeholder):
        cleaned_name = re.sub(r"[\{\}<>\[\]]+", "", placeholder).strip()
        normalized_name = re.sub(r"\s+", "_", cleaned_name.replace("'", "")).lower()
        return f"{{{{{normalized_name}}}}}" if normalized_name else placeholder

    def _replace_placeholders(self, content, placeholder_values):
        if not content:
            return ""

        soup = BeautifulSoup(content, "html.parser")
        normalized_values = {
            re.sub(r"\s+", "_", k.replace("'", "")).lower(): v
            for k, v in placeholder_values.items()
            if v is not None
        }

        placeholder_patterns = [
            r"\{\{[\w\s\'-]+\}\}",
            r"\{[\w\s\'-]+\}",
            r"\[\[[\w\s\'-]+\]\]",
            r"\[[\w\s\'-]*\w+\]",
            r"<<[\w\s\'-]+>>",
            r"<[\w\s\'-]+>",
            r"([\w\s\'-]+?)\s*:?\s*_{10,}"
        ]
        combined_pattern = "|".join(f"({pattern})" for pattern in placeholder_patterns)

        for text_node in soup.find_all(string=True):
            if not text_node.strip():
                continue
            text = text_node
            matches = re.findall(combined_pattern, text, re.IGNORECASE)
            matches = [match for group in matches for match in group if match]
            if not matches:
                continue

            new_text = text
            for match in matches:
                if re.match(r"([\w\s\'-]+?)\s*:?\s*_{10,}", match, re.IGNORECASE):
                    phrase = re.match(r"([\w\s\'-]+?)\s*:?\s*_{10,}", match, re.IGNORECASE).group(1).strip()
                    normalized_key = re.sub(r"\s+", "_", phrase.replace("'", "")).lower()
                    value = normalized_values.get(normalized_key, "__________")
                    replacement = f"{phrase}: {value}"
                    new_text = re.sub(re.escape(match), replacement, new_text, count=1)
                else:
                    normalized_key = self._normalize_placeholder(match).strip("{}").lower()
                    value = normalized_values.get(normalized_key, match)
                    value = html.escape(str(value))
                    new_text = new_text.replace(match, value)

            if new_text != text:
                text_node.replace_with(new_text)

        special_keys = ["initials", "signature", "days", "state", "employee_job_title", "employee_supervisor_name"]
        for key in special_keys:
            if key in normalized_values:
                for text_node in soup.find_all(string=re.compile(rf"(?i){key}\s*:?\s*_{{10,}}")):
                    replacement = f"{key.title()}: {html.escape(str(normalized_values[key]))}"
                    text_node.replace_with(text_node.replace(
                        re.compile(rf"(?i){key}\s*:?\s*_{{10,}}"),
                        replacement
                    ))

        return str(soup).strip()

    def generate_word_document(self, template, placeholder_values):
        if not template.raw_content or template.template_type != "word":
            return None

        doc = DocxDocument()
        normalized_values = {
            re.sub(r"\s+", "_", k.replace("'", "")).lower(): v
            for k, v in placeholder_values.items()
            if v is not None
        }

        paragraphs = template.raw_content.split("\n")
        for para_text in paragraphs:
            if not para_text.strip():
                continue
            new_text = para_text
            for placeholder in template.placeholders:
                normalized_key = re.sub(r"[\{\}<>\[\]]+", "", placeholder).strip().replace(" ", "_").lower()
                value = normalized_values.get(normalized_key, placeholder)
                new_text = new_text.replace(placeholder, str(value))
            para = doc.add_paragraph()
            if para_text.strip().startswith(("•", "◦", "-")):
                para.style = "List Bullet"
            elif re.match(r"^\d+\.", para_text.strip()):
                para.style = "List Number"
            elif para_text.strip().isupper():
                para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = para.add_run(new_text)
                run.bold = True
                run.font.size = Pt(14)
            else:
                para.add_run(new_text)

        temp_file = f"temp_{template.id}.docx"
        doc.save(temp_file)
        return temp_file

class GenerateDocumentView(BaseDocumentView):
    def get(self, request, template_id):
        template = get_object_or_404(DocumentTemplate, pk=template_id)
        template_placeholders = template.placeholders or []

        system_config = SystemConfiguration.objects.filter(
            code="doc_required_fields"
        ).first()
        required_placeholders = system_config.content if system_config else []
        required_placeholders = [
            ph for ph in required_placeholders if ph in template_placeholders
        ]

        clean_template_placeholders = [p.strip("{}").lower() for p in template_placeholders]
        all_placeholders = list(set(clean_template_placeholders + required_placeholders))

        context = request.query_params.get("context")
        context_id = request.query_params.get("context_id")
        if not context or not context_id:
            return Response(
                {"error": "context and context_id are required"},
                status=status.HTTP_400_BAD_REQUEST,
            )

        known_values = {}
        if context == "onboarding":
            onboarding = get_object_or_404(OnBoarding, pk=context_id)
            known_values = {
                "fullname": (
                    onboarding.application.applicant_name
                    if onboarding.application
                    else ""
                ),
                "salary": (
                    str(onboarding.application.job_position_advert.job_position.salary_min)
                    if onboarding.application.job_position_advert.job_position.salary_min
                    else ""
                ),
                "date": str(datetime.now().date()),
                "employee_job_title": (
                    onboarding.application.job_position_advert.job_position.name
                    if onboarding.application and onboarding.application.job_position_advert
                    else ""
                )
            }
        elif context == "employee":
            employee = get_object_or_404(Employee, pk=context_id)
            known_values = {
                "fullname": (
                    employee.user.fullname if hasattr(employee.user, "fullname") else ""
                ),
                "salary": str(employee.salary) if hasattr(employee, "salary") else "",
                "date": str(datetime.now().date()),
                "employee_job_title": (
                    employee.position.name if hasattr(employee, "position") else ""
                )
            }
        elif context == "leave":
            leave = get_object_or_404(Leave, pk=context_id)
            known_values = {
                "fullname": (
                    leave.employee.fullname
                    if hasattr(leave.employee, "fullname")
                    else ""
                ),
                "date": str(datetime.now().date()),
            }
        else:
            return Response(
                {"error": "Invalid context"}, status=status.HTTP_400_BAD_REQUEST
            )

        placeholder_data = {ph: {"value": known_values.get(ph, "")} for ph in all_placeholders}
        serializer = GenerateDocumentResponseSerializer(
            {"placeholders": placeholder_data, "template_id": template_id}
        )
        return Response(serializer.data)

    def post(self, request, template_id):
        template = get_object_or_404(DocumentTemplate, pk=template_id)
        serializer = GenerateDocumentRequestSerializer(data=request.data)
        if not serializer.is_valid():
            return Response(serializer.errors, status=status.HTTP_400_BAD_REQUEST)

        placeholder_values = serializer.validated_data["placeholders"]
        context = serializer.validated_data.get("context")
        context_id = serializer.validated_data.get("context_id")

        html_content = self._replace_placeholders(template.content, placeholder_values)
        word_file = None
        if template.template_type == "word":
            word_file = self.generate_word_document(template, placeholder_values)

        document = Document.objects.create(
            document_template=template,
            placeholder_values=placeholder_values,
            content=html_content,
            status="pending",
        )

        if context == "onboarding" and context_id:
            try:
                onboarding = OnBoarding.objects.get(pk=context_id)
                onboarding.status = "contract_review"
                onboarding.save()
            except OnBoarding.DoesNotExist:
                return Response(
                    {"error": "OnBoarding record not found"},
                    status=status.HTTP_404_NOT_FOUND,
                )

        response_data = {
            "status": "success",
            "document_id": document.pk,
            "html_content": html_content,
        }
        if word_file:
            response_data["word_file"] = word_file

        return Response(response_data, status=status.HTTP_201_CREATED)

class DocumentContentPreviewView(BaseDocumentView):
    @extend_schema(
        tags=["Document Generation"],
        responses={
            200: DocumentContentPreviewSerializer,
            404: {"description": "Document not found"},
        },
        description="Generates a preview of the document content by replacing placeholders with stored values",
    )
    def get(self, request, document_id):
        document = get_object_or_404(Document, pk=document_id)
        template = document.document_template
        template_content = template.content or ""
        placeholder_values = document.placeholder_values or {}

        preview = self._replace_placeholders(template_content, placeholder_values)
        serializer = DocumentContentPreviewSerializer({"preview": preview})
        return Response(serializer.data, status=status.HTTP_200_OK)

class DocumentStatusUpdateView(BaseDocumentView):
    def _generate_pdf(self, content, placeholder_values):
        logger.debug(f"Generating PDF with content length: {len(content)}")
        rendered_content = self._replace_placeholders(content, placeholder_values)

        html_content = f"""
        <!DOCTYPE html>
        <html>
        <head>
            <meta charset="utf-8" />
            <title>Document</title>
            <style>
                body {{
                    margin: 1in;
                }}
                h1, h2, h3, h4, h5, h6, p, li, span {{
                    margin: 0.5em 0;
                }}
                ul, ol {{
                    margin-left: 1.5em;
                }}
            </style>
        </head>
        <body>
            {rendered_content}
        </body>
        </html>
        """

        try:
            pdf_file = HTML(string=html_content).write_pdf()
            logger.info("PDF generated successfully")
            return pdf_file
        except Exception as e:
            logger.error(f"Error generating PDF: {str(e)}")
            raise

    def _get_email_values(self, placeholder_values, context, context_obj):
        email_values = {
            "due_date": str(datetime.now().date() + timedelta(days=7)),
            "hr_email": getattr(settings, "HR_EMAIL", settings.DEFAULT_FROM_EMAIL),
            "word_attachment": context_obj.document_template.template_type == "word" if hasattr(context_obj, "document_template") else False,
        }

        if context == "onboarding":
            job_position = (
                context_obj.application.job_position_advert.job_position
                if context_obj.application and context_obj.application.job_position_advert
                else None
            )
            email_values.update(
                {
                    "employee_name": (
                        context_obj.application.applicant_name
                        if context_obj.application
                        else ""
                    ),
                    "position_title": job_position.name if job_position else "",
                    "institution_name": (
                        context_obj.application.job_position_advert.job_position.department.institution.institution_name
                        if context_obj.application
                        and context_obj.application.job_position_advert
                        and context_obj.application.job_position_advert.job_position
                        and context_obj.application.job_position_advert.job_position.department
                        else ""
                    ),
                }
            )
        elif context == "employee":
            email_values.update(
                {
                    "employee_name": (
                        context_obj.user.fullname if hasattr(context_obj, "user") else ""
                    ),
                    "position_title": (
                        context_obj.job_position.name
                        if hasattr(context_obj, "job_position")
                        else ""
                    ),
                }
            )
        elif context == "leave":
            email_values.update(
                {
                    "employee_name": (
                        context_obj.employee.user.fullname
                        if hasattr(context_obj.employee, "user")
                        else ""
                    ),
                    "leave_type": (
                        context_obj.leave_type if hasattr(context_obj, "leave_type") else ""
                    ),
                    "start_date": (
                        str(context_obj.start_date) if hasattr(context_obj, "start_date") else ""
                    ),
                    "end_date": (
                        str(context_obj.end_date) if hasattr(context_obj, "end_date") else ""
                    ),
                }
            )

        for key in email_values:
            if key in placeholder_values:
                email_values[key] = placeholder_values[key]
        logger.debug(f"Email values for {context}: {email_values}")
        return email_values

    def _get_context_object(self, context, context_id):
        if context == "onboarding":
            return get_object_or_404(OnBoarding, pk=context_id)
        elif context == "employee":
            return get_object_or_404(Employee, pk=context_id)
        elif context == "leave":
            return get_object_or_404(Leave, pk=context_id)
        raise ValueError("Invalid context")

    def patch(self, request, document_id):
        context = request.query_params.get("context")
        context_id = request.query_params.get("context_id")
        if not context or not context_id:
            return Response(
                {"error": "context and context_id are required"},
                status=status.HTTP_400_BAD_REQUEST,
            )

        document = get_object_or_404(Document, pk=document_id)
        serializer = DocumentStatusUpdateSerializer(data=request.data)
        if not serializer.is_valid():
            return Response(serializer.errors, status=status.HTTP_400_BAD_REQUEST)

        new_status = serializer.validated_data["status"]
        placeholder_values = document.placeholder_values or {}

        document.status = new_status
        document.save()

        try:
            context_obj = self._get_context_object(context, context_id)
        except ValueError:
            return Response(
                {"error": "Invalid context"}, status=status.HTTP_400_BAD_REQUEST
            )
        except Exception as e:
            return Response(
                {"error": f"Context record not found: {str(e)}"},
                status=status.HTTP_404_NOT_FOUND,
            )

        template = document.document_template
        template_content = template.content or ""

        if (
            context == "onboarding"
            and hasattr(context_obj, "application")
            and context_obj.application.job_position_advert
        ):
            job_position = context_obj.application.job_position_advert.job_position
            template_content = (
                job_position.contract_template.content
                if job_position and job_position.contract_template
                else template.content
            )
        elif (
            context == "employee"
            and hasattr(context_obj, "job_position")
        ):
            template_content = (
                context_obj.job_position.contract_template.content
                if context_obj.job_position and context_obj.job_position.contract_template
                else template.content
            )

        if new_status == "reviewed":
            try:
                email_values = self._get_email_values(
                    placeholder_values, context, context_obj
                )
                pdf_content = self._generate_pdf(template_content, placeholder_values)
                pdf_filename = f"document_{context}_{document.pk}.pdf"
                document_file = ContentFile(pdf_content, name=pdf_filename)

                word_file = None
                word_filename = None
                if template.template_type == "word":
                    word_file = self.generate_word_document(template, placeholder_values)
                    word_filename = f"document_{context}_{document.pk}.docx"

                if context == "onboarding":
                    employee_contract = EmployeeContract.objects.create(
                        applicant=context_obj.application,
                        is_active=False,
                        original_contract=document_file,
                    )
                    context_obj.status = "issued_contract"
                    context_obj.save()
                elif context == "employee":
                    employee_contract = EmployeeContract.objects.create(
                        employee=context_obj,
                        is_active=False,
                        original_contract=document_file,
                    )
                    context_obj.status = "issued_contract"
                    context_obj.save()
                elif context == "leave":
                    pass

                recipient_email = None
                if context == "onboarding" and hasattr(context_obj, "application"):
                    recipient_email = getattr(
                        context_obj.application, "applicant_email", None
                    )
                elif context == "employee" and hasattr(context_obj, "email"):
                    recipient_email = context_obj.email
                elif context == "leave" and hasattr(context_obj.employee, "email"):
                    recipient_email = context_obj.employee.email

                if recipient_email:
                    try:
                        html_message = render_to_string(
                            "emails/contract_email.html", email_values
                        )
                        text_message = render_to_string(
                            "emails/contract_email.txt", email_values
                        )
                        email = EmailMultiAlternatives(
                            subject=f"Your {context.capitalize()} Document",
                            body=text_message,
                            from_email=settings.DEFAULT_FROM_EMAIL,
                            to=[recipient_email],
                        )
                        email.attach_alternative(html_message, "text/html")
                        email.attach(pdf_filename, pdf_content, "application/pdf")
                        if word_file:
                            with open(word_file, "rb") as f:
                                email.attach(word_filename, f.read(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
                        email.send(fail_silently=False)
                        logger.info(f"Email sent to {recipient_email}")
                    except Exception as e:
                        logger.error(f"Error sending {context} email: {str(e)}")
                        raise

            except Exception as e:
                logger.error(f"Error processing {context} document: {str(e)}")
                return Response(
                    {"error": f"Error processing document: {str(e)}"},
                    status=status.HTTP_500_INTERNAL_SERVER_ERROR,
                )

        preview = self._replace_placeholders(template_content, placeholder_values)
        serializer = DocumentContentPreviewSerializer({"preview": preview})
        return Response(serializer.data, status=status.HTTP_200_OK)
